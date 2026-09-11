"""Kabelbed — Fase 1-prototype (FO §9).

FastAPI-backend: datalagen ophalen, tracé rekenen, registers vullen,
exporteren. Serveert ook de frontend (map ../frontend) op /.
"""
from __future__ import annotations

import io
import json
import math
import os
import re
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from shapely.geometry import LineString, MultiPoint, Point, Polygon, mapping
from shapely.ops import unary_union

import nota as nota_mod
import pdok
import zro as zro_mod
from calculatie import build_raw_calculatie
from engine import (
    CLASS_NAMES, EngineError, Grid, ZONE_NAMES, beoordeel_werkterreinen,
    build_painter, build_segments,
    detect_crossings, propose_moffen, route_chunk, shortest_path,
    straighten_iteratief, zone_lengtes, _free_station, _techniek_voor_kruising,
    DEFAULT_WEIGHTS, BOOM_WORTELZONE_M,
)
from registers import (
    VARIANT_PROFIELEN, build_boringen, build_checks, build_kosten,
    build_mca_row, build_onderzoeken, build_planning, build_vergunningen,
    build_werkpakketten, build_zro, ken_werkpakketten_toe,
)

ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT / "data"
DATA_DIR.mkdir(exist_ok=True)

MAX_GEBIED_KM2 = 3.0
BBOX_BUFFER_M = 25.0
AUTO_GEBIED_BUFFER_M = 150.0  # zoekruimte rond de stations als er geen gebied is getekend

# Lange tracés (corridor-modus): boven CORRIDOR_VANAF_M wordt niet één raster
# over het hele gebied gebouwd, maar per deeltraject van CHUNK_M een corridor
# van ± CORRIDOR_BREEDTE_M rond de hemelsbrede lijn tussen de stations. De
# deelroutes worden aaneengehecht; het naadpunt tussen twee deeltrajecten mag
# NAAD_RADIUS_M verschuiven naar de goedkoopst bereikbare cel.
MAX_TRACE_KM = 70.0
CORRIDOR_VANAF_M = 2500.0
CHUNK_M = 1500.0
CORRIDOR_BREEDTE_M = 160.0
NAAD_RADIUS_M = 60.0
CORRIDOR_CEL = 1.0

# Zelfherstellend zoeken: vindt een deeltraject binnen de standaardcorridor
# geen doorgang (Natura 2000, aaneengesloten bebouwing, verboden zone), dan
# wordt de corridor stapsgewijs verbreed en opnieuw gerekend voordat de
# berekening opgeeft.
BREEDTE_ESCALATIE = (CORRIDOR_BREEDTE_M, 2 * CORRIDOR_BREEDTE_M,
                     4 * CORRIDOR_BREEDTE_M, 8 * CORRIDOR_BREEDTE_M)

app = FastAPI(title="Kabelbed", version="0.1")

LAST_RESULT: dict | None = None  # voor exports (single-user prototype)

# Voortgang van de lopende berekening (single-user prototype); de frontend
# pollt GET /api/progress zolang een lange berekening loopt.
PROGRESS: dict = {"actief": False}


def _voortgang(stap: str) -> None:
    PROGRESS["stap"] = stap


class ComputeRequest(BaseModel):
    area: list = Field(default_factory=list,
                       description="Projectgebied: [[x,y],...] in RD; leeg = afleiden uit stations")
    stations: list = Field(..., description="MS-stations in volgorde: [[x,y],...]")
    via: list = Field(default_factory=list, description="Verplichte passeerpunten")
    forbidden: list = Field(default_factory=list, description="Verboden zones: [[[x,y],...]]")
    weights: dict = Field(default_factory=dict, description="Wegingsprofiel-overrides")
    variants: bool = Field(default=False, description="Ook de drie varianten rekenen")
    haspel_m: float = Field(default=500.0)


def _or_masks(*masks):
    """Booleaanse OR van WMS-maskers; None-maskers (falende dienst) tellen niet mee."""
    result = None
    for m in masks:
        if m is None:
            continue
        result = m if result is None else (result | m)
    return result


def _insert_via(stations: list, via: list) -> list:
    """Via-punten toewijzen aan de dichtstbijzijnde verbinding, in looprichting."""
    if not via:
        return [tuple(s) for s in stations]
    legs = [[tuple(a), tuple(b)] for a, b in zip(stations[:-1], stations[1:])]
    per_leg: dict = {i: [] for i in range(len(legs))}
    for v in via:
        v = tuple(v)
        best_i = min(
            range(len(legs)),
            key=lambda i: LineString(legs[i]).distance(Point(v)),
        )
        per_leg[best_i].append(v)
    waypoints = [legs[0][0]]
    for i, leg in enumerate(legs):
        line = LineString(leg)
        vs = sorted(per_leg[i], key=lambda p: line.project(Point(p)))
        waypoints.extend(vs)
        waypoints.append(leg[1])
    return waypoints


def _fetch_lagen(bbox: tuple, cell: float, gemeente_punt: tuple | None = None) -> dict:
    """Alle datalagen voor één werkgebied (bbox), parallel opgehaald.

    Gedeeld door de gewone berekening en de corridor-modus (per deeltraject).
    Zonelagen zijn tolerant: een falende dienst betekent minder scherpte, geen
    fout (gemeld in laag_fouten).
    """
    fouten: list = []

    def veilig(naam, fut, leeg):
        try:
            return fut.result()
        except Exception as e:
            fouten.append(f"{naam}: {e}")
            return leeg

    dims = Grid(bbox, cell)  # alleen voor rasterafmetingen van de WMS-maskers
    cx, cy = gemeente_punt or ((bbox[0] + bbox[2]) / 2, (bbox[1] + bbox[3]) / 2)
    with ThreadPoolExecutor(max_workers=10) as ex:
        mask = lambda url, laag: ex.submit(pdok.fetch_wms_mask, url, laag,
                                           bbox, dims.ncols, dims.nrows)
        f_bgt = ex.submit(pdok.fetch_bgt, bbox)
        f_perc = ex.submit(pdok.fetch_percelen, bbox)
        f_gem = ex.submit(pdok.gemeente_naam, cx, cy)
        f_natura = ex.submit(pdok.fetch_natura2000, bbox)
        f_amk = ex.submit(pdok.fetch_amk, bbox)
        f_bomen = ex.submit(pdok.fetch_bomen, bbox)
        f_bomen_reg = ex.submit(pdok.fetch_bomen_regionaal, bbox)
        f_reg = ex.submit(pdok.regionale_bodem_masks, bbox, dims.ncols, dims.nrows)
        f_nnn = mask(pdok.NNN_WMS, "PS.ProtectedSite")
        f_gwb = mask(pdok.GWB_WMS, "AM.DrinkingWaterProtectionArea")
        f_sld1 = mask(pdok.SLD_WMS, "sld_contaminated_area")
        f_sld2 = mask(pdok.SLD_WMS, "sld_aftercare_area")
        f_sad = mask(pdok.SAD_WMS, "sad")
        f_wbb = mask(pdok.BODEMLOKET_WMS, "WBB_locaties")
        f_dek = mask(pdok.BODEMLOKET_WMS, "Beschikbaarheid_gegevens")

        bgt = veilig("bgt", f_bgt, {})
        percelen = veilig("dkk-percelen", f_perc, [])
        gemeente = veilig("gemeente", f_gem, None)
        natura = veilig("natura2000", f_natura, [])
        amk = veilig("amk", f_amk, [])
        bomen_bgt = veilig("bomen (BGT vegetatieobject)", f_bomen, [])
        bomen_reg, bomen_bronnen = veilig("bomen (regionaal)", f_bomen_reg, ([], []))
        reg_verontreinigd, reg_onderzoek, bodem_bronnen = veilig(
            "bodem regionaal", f_reg, ([], [], []))
        nnn_mask, gwb_mask = f_nnn.result(), f_gwb.result()
        sld1, sld2 = f_sld1.result(), f_sld2.result()
        sad, wbb, dek = f_sad.result(), f_wbb.result(), f_dek.result()

    # bomen samenvoegen met ontdubbeling (~2 m raster): dezelfde boom kan in de
    # BGT én een gemeentelijk register staan; het register gaat voor (kroon)
    bomen = []
    bomen_gezien = set()
    for g, p in bomen_reg + bomen_bgt:
        sleutel = (round(g.x / 2), round(g.y / 2))
        if sleutel in bomen_gezien:
            continue
        bomen_gezien.add(sleutel)
        bomen.append((g, p))
    boom_zones = [(g, max(BOOM_WORTELZONE_M, (p.get("kroon_m") or 0.0) / 2.0))
                  for g, p in bomen]

    zone_data = {
        "natura": [g for g, _ in natura],
        "archeo": [g for g, _ in amk],
        "bomen": [g.buffer(r) for g, r in boom_zones],
        "nnn_mask": nnn_mask,
        "gwb_mask": gwb_mask,
        "bodem_mask": _or_masks(sld1, sld2, *reg_verontreinigd),
        "bodem_onderzoek_mask": _or_masks(sad, wbb, *reg_onderzoek),
        "bodem_dekking_mask": dek,
    }
    for naam, sleutel in (("natuurnetwerk", "nnn_mask"), ("grondwaterbescherming", "gwb_mask"),
                          ("bodem BRO SLD", "bodem_mask"),
                          ("bodem BRO SAD/Wbb", "bodem_onderzoek_mask")):
        if zone_data[sleutel] is None:
            fouten.append(f"{naam}: WMS-masker niet beschikbaar")
    return {
        "bgt": bgt, "percelen": percelen, "gemeente": gemeente,
        "zone_data": zone_data, "boom_zones": boom_zones,
        "bomen_bronnen": ((["BGT vegetatieobject"] if bomen_bgt else [])
                          + list(bomen_bronnen)),
        "bodem_bronnen": list(bodem_bronnen),
        "laag_fouten": fouten,
    }


# ---------------------------------------------------------------------------
# Corridor-modus: lange tracés per deeltraject (tot MAX_TRACE_KM)
# ---------------------------------------------------------------------------

def _corridor_chunks(waypoints: list) -> list:
    """Hemelsbrede lijn opdelen in deeltrajecten van ± CHUNK_M.

    Deeltrajecten breken nooit door een station of via-punt heen: elke
    verbinding wordt apart opgedeeld. `vast` markeert een eindpunt dat een
    echt station/via-punt is (naad ligt daar vast)."""
    chunks = []
    for a, b in zip(waypoints[:-1], waypoints[1:]):
        d = math.hypot(b[0] - a[0], b[1] - a[1])
        if d < 1.0:
            continue
        n = max(1, math.ceil(d / CHUNK_M))
        for i in range(n):
            t0, t1 = i / n, (i + 1) / n
            chunks.append({
                "van": (a[0] + t0 * (b[0] - a[0]), a[1] + t0 * (b[1] - a[1])),
                "tot": (a[0] + t1 * (b[0] - a[0]), a[1] + t1 * (b[1] - a[1])),
                "vast": i == n - 1,
            })
    return chunks


def _chunk_bbox(chunk: dict, breedte: float = CORRIDOR_BREEDTE_M) -> tuple:
    b = LineString([chunk["van"], chunk["tot"]]).buffer(breedte).bounds
    return (b[0] - BBOX_BUFFER_M, b[1] - BBOX_BUFFER_M,
            b[2] + BBOX_BUFFER_M, b[3] + BBOX_BUFFER_M)


def _hecht_kruisingen(kruisingen: list, route: LineString) -> list:
    """Kruisingen van alle deeltrajecten samenvoegen en hernummeren.

    Een kruising die precies op een naad tussen twee deeltrajecten valt is in
    beide delen half gedetecteerd; delen van dezelfde soort die elkaar binnen
    5 m raken worden samengevoegd en krijgen een nieuw techniekvoorstel voor
    de samengevoegde breedte (de werkterrein-toets vervalt dan: het raster
    van het deeltraject is al opgeruimd)."""
    kruisingen.sort(key=lambda c: (c["soort"], c["chainage_van_m"]))
    samengevoegd: list = []
    for c in kruisingen:
        v = samengevoegd[-1] if samengevoegd else None
        if (v and v["soort"] == c["soort"]
                and c["chainage_van_m"] - v["chainage_tot_m"] < 5.0):
            v["chainage_tot_m"] = max(v["chainage_tot_m"], c["chainage_tot_m"])
            v["kruislengte_m"] = round(v["chainage_tot_m"] - v["chainage_van_m"], 1)
            # haakse breedte: elk half gedetecteerd deel mat al door het hele
            # obstakel; de grootste meting is de beste schatting
            v["breedte_m"] = max(v["breedte_m"], c["breedte_m"])
            v.update(_techniek_voor_kruising(v["soort"], v["breedte_m"]))
            v.pop("werkterrein", None)
            v.pop("techniek_oorspronkelijk", None)
            mid = route.interpolate((v["chainage_van_m"] + v["chainage_tot_m"]) / 2)
            v["punt"] = (round(mid.x, 2), round(mid.y, 2))
        else:
            samengevoegd.append(c)
    samengevoegd.sort(key=lambda c: c["chainage_van_m"])
    for i, c in enumerate(samengevoegd, 1):
        c["nr"] = f"KR-{i:03d}"
    return samengevoegd


def _hecht_segmenten(segmenten: list) -> list:
    """Aansluitende segmenten met dezelfde ligging over naden heen samenvoegen.

    De drempel is de sample-afstand van build_segments (2 m): opeenvolgende
    runs liggen daardoor maximaal één sample uit elkaar, ook op een naad."""
    out: list = []
    for s in segmenten:  # al in chainage-volgorde
        v = out[-1] if out else None
        if v and v["klasse"] == s["klasse"] and s["van_m"] - v["tot_m"] < 2.1:
            v["tot_m"] = s["tot_m"]
            v["lengte_m"] = round(v["tot_m"] - v["van_m"], 1)
            v["geometry"] = {"type": "LineString",
                             "coordinates": (list(v["geometry"]["coordinates"])
                                             + list(s["geometry"]["coordinates"])[1:])}
        else:
            out.append(s)
    for i, s in enumerate(out, 1):
        s["nr"] = f"SEG-{i:03d}"
    return out


def _compute_corridor(req: ComputeRequest, waypoints: list, hemelsbreed: float,
                      t0: float) -> dict:
    """Lang tracé: per deeltraject datalagen ophalen, corridor-raster bouwen
    en routeren; daarna alles aaneenhechten en de registers op het volledige
    tracé vullen. Datalagen van volgende deeltrajecten worden vooruit
    opgehaald terwijl het huidige deeltraject rekent."""
    chunks = _corridor_chunks(waypoints)
    profielen = dict(VARIANT_PROFIELEN) if req.variants else {"Voorkeursvariant": {}}
    staat = {naam: {"coords": [], "lengte": 0.0, "kruisingen": [], "segmenten": [],
                    "zones": {bit: 0.0 for bit in ZONE_NAMES}, "fout": None}
             for naam in profielen}
    gemeenten: list = []
    bomen_alle: list = []
    bomen_gezien: set = set()
    percelen_alle: dict = {}
    laag_fouten: dict = {}
    bgt_fouten: dict = {}
    bron_namen = {"bodem": [], "bomen": []}
    t_data = 0.0

    def merk_data(data: dict) -> None:
        """Statistieken, percelen en bomen van een (extra) databundel bijhouden."""
        for f in data["laag_fouten"]:
            naam = f.split(":")[0]
            laag_fouten[naam] = laag_fouten.get(naam, 0) + 1
        for f in data["bgt"].get("_errors", []):
            naam = f.split(":")[0]
            bgt_fouten[naam] = bgt_fouten.get(naam, 0) + 1
        if data["gemeente"] and data["gemeente"] not in gemeenten:
            gemeenten.append(data["gemeente"])
        for soort, sleutel in (("bodem", "bodem_bronnen"), ("bomen", "bomen_bronnen")):
            for naam in data[sleutel]:
                if naam not in bron_namen[soort]:
                    bron_namen[soort].append(naam)
        for g, p in data["percelen"]:
            sleutel = p.get("identificatieLokaalID") or (
                p.get("kadastraleGemeenteWaarde"), p.get("sectie"),
                p.get("perceelnummer"))
            percelen_alle.setdefault(sleutel, (g, p))
        for g, r in data["boom_zones"]:
            sleutel = (round(g.x / 2), round(g.y / 2))
            if sleutel not in bomen_gezien:
                bomen_gezien.add(sleutel)
                bomen_alle.append((g, r))

    verbreed: list = []  # [{deeltraject, breedte_m}] voor de melding aan de gebruiker

    with ThreadPoolExecutor(max_workers=2) as pool:
        futs: dict = {}

        def plan(j):
            if 0 <= j < len(chunks) and j not in futs:
                futs[j] = pool.submit(_fetch_lagen, _chunk_bbox(chunks[j]), CORRIDOR_CEL)

        for j in range(3):
            plan(j)
        for i, chunk in enumerate(chunks):
            _voortgang(f"deeltraject {i + 1}/{len(chunks)}: datalagen ophalen")
            t_wacht = time.time()
            data = futs.pop(i).result()
            t_data += time.time() - t_wacht
            plan(i + 2)
            plan(i + 3)
            merk_data(data)

            # datalagen en painters per corridorbreedte, lui opgebouwd: de
            # bredere varianten worden alleen opgehaald als de standaardbreedte
            # geen doorgang oplevert (zelfherstellend zoeken)
            datas = {BREEDTE_ESCALATIE[0]: data}
            painters: dict = {}

            def painter_voor(breedte):
                if breedte not in painters:
                    if breedte not in datas:
                        _voortgang(f"deeltraject {i + 1}/{len(chunks)}: geen "
                                   f"doorgang — corridor verbreden naar "
                                   f"±{breedte:.0f} m, datalagen ophalen")
                        datas[breedte] = _fetch_lagen(
                            _chunk_bbox(chunk, breedte), CORRIDOR_CEL)
                        merk_data(datas[breedte])
                    d = datas[breedte]
                    painters[breedte] = build_painter(
                        _chunk_bbox(chunk, breedte), d["bgt"], req.forbidden,
                        CORRIDOR_CEL, d["zone_data"])
                return painters[breedte]

            for naam, overrides in profielen.items():
                st = staat[naam]
                if st["fout"]:
                    continue
                _voortgang(f"deeltraject {i + 1}/{len(chunks)}: route rekenen ({naam})")
                weights = {**DEFAULT_WEIGHTS, **req.weights, **overrides}
                start = tuple(st["coords"][-1]) if st["coords"] else tuple(waypoints[0])

                coords = grid = None
                fout = None
                for breedte in BREEDTE_ESCALATIE:
                    try:
                        grid = painter_voor(breedte).build(weights)
                        _free_station(grid, start)
                        _free_station(grid, chunk["tot"])
                        coords = route_chunk(grid, start, chunk["tot"],
                                             0.0 if chunk["vast"] else NAAD_RADIUS_M)
                        fout = None
                        if breedte != BREEDTE_ESCALATIE[0]:
                            chunk["breedte"] = max(chunk.get("breedte", 0.0), breedte)
                            verbreed.append({"deeltraject": i + 1,
                                             "breedte_m": breedte})
                        break
                    except EngineError as e:
                        fout = e
                        _voortgang(f"deeltraject {i + 1}/{len(chunks)}: geen "
                                   f"doorgang binnen ±{breedte:.0f} m ({naam})")
                if fout is not None:
                    st["fout"] = (f"Deeltraject {i + 1}/{len(chunks)}: {fout} "
                                  f"(automatisch geprobeerd tot een corridor van "
                                  f"±{BREEDTE_ESCALATIE[-1]:.0f} m)")
                    continue

                try:
                    bgt_deel = datas[breedte]["bgt"]
                    deel = LineString(coords)
                    deel = straighten_iteratief(deel, bgt_deel, grid)
                    kruisingen = detect_crossings(deel, bgt_deel)
                    beoordeel_werkterreinen(deel, kruisingen, grid)
                    segmenten = build_segments(deel, grid)
                    zl = zone_lengtes(deel, grid)
                except EngineError as e:
                    st["fout"] = f"Deeltraject {i + 1}/{len(chunks)}: {e}"
                    continue
                offset = st["lengte"]
                for c in kruisingen:
                    c["chainage_van_m"] = round(c["chainage_van_m"] + offset, 1)
                    c["chainage_tot_m"] = round(c["chainage_tot_m"] + offset, 1)
                for s in segmenten:
                    s["van_m"] = round(s["van_m"] + offset, 1)
                    s["tot_m"] = round(s["tot_m"] + offset, 1)
                st["kruisingen"].extend(kruisingen)
                st["segmenten"].extend(segmenten)
                for bit, m in zl.items():
                    st["zones"][bit] += m
                cs = list(deel.coords)
                st["coords"].extend(cs if not st["coords"] else cs[1:])
                st["lengte"] += deel.length
            del painters, datas, data

    _voortgang("registers en toetsing samenstellen")
    gemeente = " / ".join(gemeenten) if gemeenten else None
    percelen = list(percelen_alle.values())
    varianten, fouten = [], []
    for naam in profielen:
        st = staat[naam]
        if st["fout"]:
            fouten.append({"variant": naam, "fout": st["fout"]})
            continue
        route = LineString(st["coords"])
        kruisingen = _hecht_kruisingen(st["kruisingen"], route)
        segmenten = _hecht_segmenten(st["segmenten"])
        zones_m = {bit: round(m, 1) for bit, m in st["zones"].items()}
        moffen = propose_moffen(route, kruisingen, segmenten, req.haspel_m)
        vergunningen = build_vergunningen(segmenten, kruisingen, gemeente, zones_m)
        boringen = build_boringen(route, kruisingen)
        onderzoeken = build_onderzoeken(zones_m, boringen)
        zro = build_zro(route, percelen)
        checks = build_checks(route, segmenten, kruisingen, boringen, zones_m)
        werkpakketten = build_werkpakketten(route, req.stations)
        ken_werkpakketten_toe(werkpakketten, route, segmenten, kruisingen,
                              boringen, moffen, zro, vergunningen,
                              onderzoeken, checks)
        planning = build_planning(werkpakketten, vergunningen, onderzoeken,
                                  boringen, zro)
        kosten = build_kosten(segmenten, kruisingen, zro, moffen, onderzoeken)
        calculatie = build_raw_calculatie(route.length, segmenten, kruisingen,
                                          boringen, zro, moffen, onderzoeken,
                                          zones_m, kosten, len(req.stations))
        kosten["aannemingssom_excl_btw"] = calculatie["aannemingssom_excl_btw"]
        mca = build_mca_row(naam, route, segmenten, kruisingen, zro, vergunningen, kosten)
        varianten.append({
            "naam": naam,
            "wegingsprofiel": {**DEFAULT_WEIGHTS, **req.weights, **profielen[naam]},
            "route": mapping(route),
            "lengte_m": round(route.length, 1),
            "kruisingen": kruisingen,
            "segmenten": segmenten,
            "zones": {ZONE_NAMES[bit]: m for bit, m in zones_m.items() if m > 0},
            "moffen": moffen,
            "vergunningen": vergunningen,
            "boringen": boringen,
            "onderzoeken": onderzoeken,
            "zro": zro,
            "toetsing": checks,
            "werkpakketten": werkpakketten,
            "planning": planning,
            "kosten": kosten,
            "calculatie": calculatie,
            "mca": mca,
        })

    if not varianten:
        raise HTTPException(422, fouten[0]["fout"] if fouten else "Berekening mislukt.")

    corridor = unary_union(
        [LineString([c["van"], c["tot"]]).buffer(c.get("breedte", CORRIDOR_BREEDTE_M))
         for c in chunks]).simplify(10)
    max_b = max((c.get("breedte", CORRIDOR_BREEDTE_M) for c in chunks),
                default=CORRIDOR_BREEDTE_M)
    # per deeltraject de grootste gebruikte verbreding melden
    verbreed_per_deel: dict = {}
    for v in verbreed:
        verbreed_per_deel[v["deeltraject"]] = max(
            verbreed_per_deel.get(v["deeltraject"], 0), v["breedte_m"])
    xs = [p[0] for c in chunks for p in (c["van"], c["tot"])]
    ys = [p[1] for c in chunks for p in (c["van"], c["tot"])]
    return {
        "modus": "corridor",
        "gemeente": gemeente,
        "celgrootte_m": CORRIDOR_CEL,
        "corridor_verbreed": [{"deeltraject": d, "breedte_m": b}
                              for d, b in sorted(verbreed_per_deel.items())],
        "bbox": (min(xs) - max_b, min(ys) - max_b,
                 max(xs) + max_b, max(ys) + max_b),
        "hemelsbreed_m": round(hemelsbreed),
        "deeltrajecten": len(chunks),
        "rekentijd_s": {"datalagen": round(t_data, 1),
                        "route": round(time.time() - t0 - t_data, 1)},
        "bgt_fouten": [f"{naam} ({n}× deeltraject)" for naam, n in sorted(bgt_fouten.items())],
        "laag_fouten": [f"{naam} ({n}× deeltraject)" for naam, n in sorted(laag_fouten.items())],
        "bodem_bronnen_regionaal": bron_namen["bodem"],
        "gebied": mapping(corridor),
        "bomen": [[round(g.x, 2), round(g.y, 2), round(r, 1)] for g, r in bomen_alle],
        "bomen_bronnen": bron_namen["bomen"],
        "varianten": varianten,
        "variant_fouten": fouten,
        "stations": req.stations,
    }


@app.post("/api/compute")
def compute(req: ComputeRequest):
    global LAST_RESULT
    t0 = time.time()
    if len(req.stations) < 2:
        raise HTTPException(400, "Plaats minimaal twee MS-stations.")

    waypoints = _insert_via(req.stations, req.via)
    hemelsbreed = sum(math.hypot(b[0] - a[0], b[1] - a[1])
                      for a, b in zip(waypoints[:-1], waypoints[1:]))
    if hemelsbreed > MAX_TRACE_KM * 1000:
        raise HTTPException(400, f"Tracé is hemelsbreed {hemelsbreed / 1000:.1f} km; "
                                 f"maximum voor dit prototype is {MAX_TRACE_KM:.0f} km.")

    PROGRESS.update({"actief": True, "stap": "voorbereiden", "t0": t0})
    try:
        if hemelsbreed > CORRIDOR_VANAF_M:
            # lang tracé: per deeltraject rekenen (een eventueel getekend
            # projectgebied wordt genegeerd; verboden zones blijven gelden)
            result = _compute_corridor(req, waypoints, hemelsbreed, t0)
            LAST_RESULT = result
            return result
        result = _compute_gebied(req, waypoints, t0)
        LAST_RESULT = result
        return result
    finally:
        PROGRESS["actief"] = False


def _compute_gebied(req: ComputeRequest, waypoints: list, t0: float) -> dict:
    """Korte tracés: één raster over het (getekende of afgeleide) projectgebied."""
    if req.area:
        if len(req.area) < 3:
            raise HTTPException(400, "Projectgebied moet minimaal drie punten hebben.")
        gebied = Polygon(req.area)
        km2 = gebied.area / 1e6
        if km2 > MAX_GEBIED_KM2:
            raise HTTPException(400, f"Projectgebied is {km2:.1f} km²; maximum voor dit "
                                     f"prototype is {MAX_GEBIED_KM2} km².")
        for s in req.stations:
            if not gebied.buffer(5).contains(Point(s)):
                raise HTTPException(400, "Alle stations moeten binnen het projectgebied liggen.")
    else:
        punten = [tuple(s) for s in req.stations] + [tuple(v) for v in req.via]
        gebied = MultiPoint(punten).convex_hull.buffer(AUTO_GEBIED_BUFFER_M)
        km2 = gebied.area / 1e6
        if km2 > MAX_GEBIED_KM2:
            raise HTTPException(400, f"Automatisch zoekgebied rond de stations is {km2:.1f} km²; "
                                     f"maximum voor dit prototype is {MAX_GEBIED_KM2} km². "
                                     f"Plaats de stations dichter bij elkaar of teken zelf "
                                     f"een projectgebied.")

    xmin, ymin, xmax, ymax = gebied.bounds
    bbox = (xmin - BBOX_BUFFER_M, ymin - BBOX_BUFFER_M, xmax + BBOX_BUFFER_M, ymax + BBOX_BUFFER_M)
    cell = 0.5 if km2 <= 1.0 else 1.0

    # --- datalagen (open bronnen), parallel ---
    _voortgang("datalagen ophalen bij PDOK")
    data = _fetch_lagen(bbox, cell, (gebied.centroid.x, gebied.centroid.y))
    bgt = data["bgt"]
    percelen = data["percelen"]
    gemeente = data["gemeente"]
    laag_fouten = data["laag_fouten"]
    boom_zones = data["boom_zones"]
    t_data = time.time()

    painter = build_painter(bbox, bgt, req.forbidden, cell, data["zone_data"])

    profielen = dict(VARIANT_PROFIELEN) if req.variants else {"Voorkeursvariant": {}}
    varianten = []
    fouten = []
    for naam, overrides in profielen.items():
        try:
            _voortgang(f"route rekenen ({naam})")
            weights = {**DEFAULT_WEIGHTS, **req.weights, **overrides}
            grid = painter.build(weights)
            for wp in waypoints:
                _free_station(grid, wp)
            route = LineString(shortest_path(grid, waypoints))
            route = straighten_iteratief(route, bgt, grid)
            crossings = detect_crossings(route, bgt)
            beoordeel_werkterreinen(route, crossings, grid)
            segments = build_segments(route, grid)
            zones_m = zone_lengtes(route, grid)
            moffen = propose_moffen(route, crossings, segments, req.haspel_m)
            vergunningen = build_vergunningen(segments, crossings, gemeente, zones_m)
            boringen = build_boringen(route, crossings)
            onderzoeken = build_onderzoeken(zones_m, boringen)
            zro = build_zro(route, percelen)
            checks = build_checks(route, segments, crossings, boringen, zones_m)
            werkpakketten = build_werkpakketten(route, req.stations)
            ken_werkpakketten_toe(werkpakketten, route, segments, crossings,
                                  boringen, moffen, zro, vergunningen,
                                  onderzoeken, checks)
            planning = build_planning(werkpakketten, vergunningen, onderzoeken,
                                      boringen, zro)
            kosten = build_kosten(segments, crossings, zro, moffen, onderzoeken)
            calculatie = build_raw_calculatie(route.length, segments, crossings,
                                              boringen, zro, moffen, onderzoeken,
                                              zones_m, kosten, len(req.stations))
            kosten["aannemingssom_excl_btw"] = calculatie["aannemingssom_excl_btw"]
            mca = build_mca_row(naam, route, segments, crossings, zro, vergunningen, kosten)
            varianten.append({
                "naam": naam,
                "wegingsprofiel": weights,
                "route": mapping(route),
                "lengte_m": round(route.length, 1),
                "kruisingen": crossings,
                "segmenten": segments,
                "zones": {ZONE_NAMES[bit]: m for bit, m in zones_m.items() if m > 0},
                "moffen": moffen,
                "vergunningen": vergunningen,
                "boringen": boringen,
                "onderzoeken": onderzoeken,
                "zro": zro,
                "toetsing": checks,
                "werkpakketten": werkpakketten,
                "planning": planning,
                "kosten": kosten,
                "calculatie": calculatie,
                "mca": mca,
            })
        except EngineError as e:
            fouten.append({"variant": naam, "fout": str(e)})

    if not varianten:
        raise HTTPException(422, fouten[0]["fout"] if fouten else "Berekening mislukt.")

    return {
        "modus": "gebied",
        "gemeente": gemeente,
        "celgrootte_m": cell,
        "bbox": bbox,
        "rekentijd_s": {"datalagen": round(t_data - t0, 1),
                        "route": round(time.time() - t_data, 1)},
        "bgt_fouten": bgt.get("_errors", []),
        "laag_fouten": laag_fouten,
        "bodem_bronnen_regionaal": data["bodem_bronnen"],
        "gebied": mapping(gebied),
        "bomen": [[round(g.x, 2), round(g.y, 2), round(r, 1)] for g, r in boom_zones],
        "bomen_bronnen": data["bomen_bronnen"],
        "varianten": varianten,
        "variant_fouten": fouten,
        "stations": req.stations,
    }


# ---------------------------------------------------------------------------
# Registers: ophalen en bewerken (FO §6 — registerpagina)
# ---------------------------------------------------------------------------

def _need_result(variant: int) -> dict:
    if LAST_RESULT is None:
        raise HTTPException(404, "Nog geen berekend tracé; reken eerst een tracé.")
    try:
        return LAST_RESULT["varianten"][variant]
    except IndexError:
        raise HTTPException(404, "Variant bestaat niet.")


# Whitelist van velden die per register vanuit de registerpagina bewerkt mogen
# worden. Afgeleide waarden (lengtes, geometrie, kosten) blijven rekenwerk.
# ZRO ontbreekt bewust: die velden lopen via het dossier (/api/zro/detail).
BEWERKBARE_VELDEN = {
    "vergunningen": {"item", "bevoegd_gezag", "trigger", "status", "verantwoordelijke"},
    "kruisingen": {"techniek", "detail", "bevoegd_gezag", "noodzaak"},
    "boringen": {"type", "mantelbuis", "dekking_eis", "noodzaak", "status",
                 "verantwoordelijke"},
    "onderzoeken": {"soort", "aanleiding", "conclusie", "status", "verantwoordelijke"},
    "moffen": {"opmerking"},
    "werkpakketten": {"naam"},
    "planning": {"status", "toelichting"},
}


class RegisterUpdate(BaseModel):
    variant: int = 0
    register: str
    nr: str | int
    wijzigingen: dict


@app.get("/api/result")
def get_result():
    if LAST_RESULT is None:
        raise HTTPException(404, "Nog geen berekend tracé.")
    return LAST_RESULT


@app.get("/api/progress")
def get_progress():
    """Voortgang van de lopende berekening (frontend pollt bij lange tracés)."""
    if PROGRESS.get("actief"):
        return {"actief": True, "stap": PROGRESS.get("stap", ""),
                "bezig_s": round(time.time() - PROGRESS.get("t0", time.time()), 1)}
    return {"actief": False}


@app.post("/api/register/update")
def register_update(u: RegisterUpdate):
    v = _need_result(u.variant)
    velden = BEWERKBARE_VELDEN.get(u.register)
    if velden is None:
        raise HTTPException(400, f"Register '{u.register}' is niet bewerkbaar.")
    rij = next((r for r in v.get(u.register, []) if str(r.get("nr")) == str(u.nr)), None)
    if rij is None:
        raise HTTPException(404, f"Rij '{u.nr}' niet gevonden in register '{u.register}'.")
    geweigerd = sorted(set(u.wijzigingen) - velden)
    if geweigerd:
        raise HTTPException(400, "Veld(en) niet bewerkbaar: " + ", ".join(geweigerd))
    for k, w in u.wijzigingen.items():
        if not isinstance(w, str) or len(w) > 500:
            raise HTTPException(400, f"Ongeldige waarde voor '{k}' (tekst, max 500 tekens).")
        rij[k] = w.strip()
    return rij


@app.get("/api/export/geojson")
def export_geojson(variant: int = 0):
    v = _need_result(variant)
    feats = [{"type": "Feature", "properties": {"laag": "trace", "variant": v["naam"],
                                                "lengte_m": v["lengte_m"]},
              "geometry": v["route"]}]
    for s in LAST_RESULT["stations"]:
        feats.append({"type": "Feature", "properties": {"laag": "station"},
                      "geometry": {"type": "Point", "coordinates": s}})
    for c in v["kruisingen"]:
        feats.append({"type": "Feature",
                      "properties": {"laag": "kruising", **{k: c.get(k) for k in
                                     ("nr", "soort", "breedte_m", "kruislengte_m",
                                      "techniek", "bevoegd_gezag", "noodzaak")},
                                     "werkpakket": c.get("werkpakket", "")},
                      "geometry": {"type": "Point", "coordinates": list(c["punt"])}})
    for b in v.get("boringen", []):
        # oudere opgeslagen projecten hebben nog geen boorlijn-geometrie
        boorlijn = b.get("geometry") or {
            "type": "LineString",
            "coordinates": [list(b["intredepunt_rd"]), list(b["uittredepunt_rd"])]}
        feats.append({"type": "Feature",
                      "properties": {"laag": "boring", **{k: b.get(k) for k in
                                     ("nr", "type", "obstakel", "kruising", "lengte_m",
                                      "uitloop_m", "dekking_eis", "noodzaak")},
                                     "werkpakket": b.get("werkpakket", "")},
                      "geometry": boorlijn})
        feats.append({"type": "Feature",
                      "properties": {"laag": "boring_intrede", "boring": b["nr"]},
                      "geometry": {"type": "Point",
                                   "coordinates": list(b["intredepunt_rd"])}})
        feats.append({"type": "Feature",
                      "properties": {"laag": "boring_uittrede", "boring": b["nr"]},
                      "geometry": {"type": "Point",
                                   "coordinates": list(b["uittredepunt_rd"])}})
    for m in v["moffen"]:
        feats.append({"type": "Feature", "properties": {"laag": "mof", "nr": m["nr"],
                                                        "werkpakket": m.get("werkpakket", "")},
                      "geometry": {"type": "Point", "coordinates": list(m["punt"])}})
    for b in LAST_RESULT.get("bomen", []):
        feats.append({"type": "Feature",
                      "properties": {"laag": "boom", "wortelzone_r_m": b[2]},
                      "geometry": {"type": "Point", "coordinates": b[:2]}})
    for s in v["segmenten"]:
        feats.append({"type": "Feature",
                      "properties": {"laag": "segment", "nr": s["nr"], "ligging": s["ligging"],
                                     "lengte_m": s["lengte_m"],
                                     "werkpakket": s.get("werkpakket", "")},
                      "geometry": s["geometry"]})
    for w in v.get("werkpakketten", []):
        feats.append({"type": "Feature",
                      "properties": {"laag": "werkpakket", "nr": w["nr"], "naam": w["naam"],
                                     "lengte_m": w["lengte_m"]},
                      "geometry": w["geometry"]})
    fc = {"type": "FeatureCollection",
          "crs": {"type": "name", "properties": {"name": "urn:ogc:def:crs:EPSG::28992"}},
          "features": feats}
    data = json.dumps(fc, ensure_ascii=False).encode()
    return StreamingResponse(io.BytesIO(data), media_type="application/geo+json",
                             headers={"Content-Disposition":
                                      f'attachment; filename="kabelbed_{variant}.geojson"'})


@app.get("/api/export/xlsx")
def export_xlsx(variant: int = 0):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    v = _need_result(variant)
    wb = Workbook()

    def sheet(title, headers, rows):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True)
        for row in rows:
            ws.append(row)
        for col in ws.columns:
            width = max((len(str(c.value or "")) for c in col), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(width + 2, 60)

    wb.remove(wb.active)
    mca_rows = []
    for var in LAST_RESULT["varianten"]:
        m = var["mca"]
        kr = "; ".join(f"{k}: {n}" for k, n in m["kruisingen"].items()) or "-"
        mca_rows.append([m["variant"], m["lengte_m"], kr, m["meters_privaat_m"],
                         m["aantal_percelen"], m["aantal_vergunningen"],
                         m["kosten_eur"], m["doorlooptijd_wk"]])
    sheet("MCA varianten",
          ["Variant", "Lengte (m)", "Kruisingen", "Privaat (m)", "Percelen",
           "Vergunningen", "Kosten (EUR)", "Doorlooptijd (wk)"], mca_rows)
    sheet("Werkpakketten",
          ["Nr", "Naam", "Van station", "Tot station", "Van (m)", "Tot (m)",
           "Lengte (m)"],
          [[w["nr"], w["naam"], w["van_station"], w["tot_station"],
            w["chainage_van_m"], w["chainage_tot_m"], w["lengte_m"]]
           for w in v.get("werkpakketten", [])])
    sheet("Planning",
          ["Nr", "Werkpakket", "Fase", "Start (wk)", "Eind (wk)", "Duur (wk)",
           "Status", "Toelichting"],
          [[p["nr"], p["werkpakket"], p["fase"], p["start_wk"], p["eind_wk"],
            p["duur_wk"], p["status"], p["toelichting"]]
           for p in v.get("planning", [])])
    sheet("Segmenten", ["Nr", "Werkpakket", "Ligging", "Van (m)", "Tot (m)", "Lengte (m)"],
          [[s["nr"], s.get("werkpakket", ""), s["ligging"], s["van_m"], s["tot_m"],
            s["lengte_m"]] for s in v["segmenten"]])
    sheet("Kruisingen",
          ["Nr", "Werkpakket", "Soort", "Breedte haaks (m)", "Langs tracé (m)",
           "Techniek", "Noodzaak", "Detail", "Richtlijn", "Bevoegd gezag"],
          [[c["nr"], c.get("werkpakket", ""), c["soort"], c["breedte_m"],
            c.get("kruislengte_m", ""), c["techniek"],
            c.get("noodzaak", ""), c["detail"], c["richtlijn"], c["bevoegd_gezag"]]
           for c in v["kruisingen"]])
    sheet("Vergunningen",
          ["Nr", "Werkpakket", "Item", "Bevoegd gezag", "Trigger", "Doorlooptijd (wk)",
           "Status", "Verantwoordelijke"],
          [[i["nr"], i.get("werkpakket", ""), i["item"], i["bevoegd_gezag"], i["trigger"],
            f"{i['doorlooptijd_wk'][0]}–{i['doorlooptijd_wk'][1]}", i["status"],
            i.get("verantwoordelijke", "")]
           for i in v["vergunningen"]])
    sheet("Boringen",
          ["Nr", "Werkpakket", "Type", "Kruising", "Obstakel", "Noodzaak",
           "Lengte (m)", "Intrede (RD)",
           "Uittrede (RD)", "Dekking-eis", "Mantelbuis",
           "Werkterrein", "Werkruimte intrede (m2)", "Werkruimte uittrede (m2)",
           "Werkterrein opmerking", "Status", "Verantwoordelijke"],
          [[b["nr"], b.get("werkpakket", ""),
            b["type"] + (f" (i.p.v. {b['type_oorspronkelijk']})"
                         if b.get("type_oorspronkelijk") else ""),
            b["kruising"], b["obstakel"], b.get("noodzaak", ""), b["lengte_m"],
            f"{b['intredepunt_rd'][0]}, {b['intredepunt_rd'][1]}",
            f"{b['uittredepunt_rd'][0]}, {b['uittredepunt_rd'][1]}",
            b["dekking_eis"], b["mantelbuis"],
            b["werkterrein_oordeel"],
            "" if b["werkterrein_intrede_m2"] is None else
            f"{b['werkterrein_intrede_m2']} (eis {b['werkterrein_intrede_eis_m2']})",
            "" if b["werkterrein_uittrede_m2"] is None else
            f"{b['werkterrein_uittrede_m2']} (eis {b['werkterrein_uittrede_eis_m2']})",
            b["werkterrein_opmerking"], b["status"],
            b.get("verantwoordelijke", "")] for b in v["boringen"]])
    zro_rows = []
    for z in v["zro"]:
        d = zro_mod.laad_dossier(zro_mod.slug_van(z["perceel"]))
        zro_rows.append([z["nr"], z.get("werkpakket", ""), z["perceel"],
                         d["eigenaar_naam"] or z["eigenaar"],
                         z["ingenomen_lengte_m"], z["werkstrook_m2"],
                         d["aard_recht"] or z["aard_recht"],
                         d["vergoeding_eenmalig_eur"],
                         d["vergoeding_jaarlijks_eur"],
                         d["vergoeding_grondslag"], d["status"],
                         "; ".join(b["bestand"] for b in d["bijlagen"])])
    sheet("ZRO",
          ["Nr", "Werkpakket", "Perceel", "Eigenaar", "Ingenomen lengte (m)",
           "Werkstrook (m2)", "Aard recht", "Vergoeding eenmalig (EUR)",
           "Vergoeding jaarlijks (EUR)", "Grondslag vergoeding", "Status",
           "Bijlagen"], zro_rows)
    sheet("Onderzoeken",
          ["Nr", "Onderzoek", "Aanleiding", "Conclusie", "Status", "Verantwoordelijke"],
          [[o["nr"], o["soort"], o["aanleiding"], o["conclusie"], o["status"],
            o.get("verantwoordelijke", "")]
           for o in v.get("onderzoeken", [])])
    sheet("Toetsing", ["Ernst", "Toets", "Grondslag", "Melding"],
          [[c["ernst"], c["toets"], c["grondslag"], c["melding"]] for c in v["toetsing"]])
    sheet("Moffen", ["Nr", "Chainage (m)", "X (RD)", "Y (RD)", "Opmerking"],
          [[m["nr"], m["chainage_m"], m["punt"][0], m["punt"][1], m["opmerking"]]
           for m in v["moffen"]])
    sheet("Kosten", ["Post", "Bedrag (EUR)"],
          [[k, w] for k, w in v["kosten"].items() if k != "toelichting"]
          + [["toelichting", v["kosten"]["toelichting"]]])
    calc = v.get("calculatie")
    if calc:
        rows = []
        for h in calc["hoofdstukken"]:
            rows.append([h["code"], "HOOFDSTUK " + h["naam"].upper(),
                         "", "", "", h["totaal_eur"], ""])
            rows.extend(
                [q["nr"], q["omschrijving"], q["eenheid"], q["hoeveelheid"],
                 q["eenheidsprijs_eur"], q["totaal_eur"], q["herkomst"]]
                for q in calc["posten"] if q["hoofdstuk"] == h["code"])
        rows.append([])
        rows.extend([s["nr"], s["omschrijving"], "", "", s["grondslag"],
                     s["bedrag_eur"], ""] for s in calc["staart"])
        rows.append([])
        rows.append(["", calc["systematiek"] + " — prijspeil " + calc["prijspeil"]])
        rows.extend(["", "• " + u] for u in calc["uitgangspunten"])
        sheet("RAW-calculatie",
              ["Bestekpost", "Omschrijving", "Eenheid", "Hoeveelheid",
               "Prijs per eenheid (EUR)", "Totaal (EUR)", "Herkomst hoeveelheid"],
              rows)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="kabelbed_registers_{variant}.xlsx"'})


DOCX_MEDIA = ("application/vnd.openxmlformats-officedocument"
              ".wordprocessingml.document")


@app.get("/api/nota/stream")
def nota_stream(fase: str, variant: int = 0, projectnaam: str = ""):
    """Ontwerpnota (VO/DO/UO) live door AI laten schrijven: streamt de
    Markdown-tekst terwijl die ontstaat, zodat de frontend het document
    ziet opbouwen. Fouten na de start staan als [NOTA-FOUT]-regel in de
    stroom."""
    _need_result(variant)
    try:
        gen = nota_mod.stream_nota(LAST_RESULT, fase, variant, projectnaam)
    except nota_mod.NotaError as e:
        raise HTTPException(503, str(e))
    return StreamingResponse(gen, media_type="text/plain; charset=utf-8",
                             headers={"Cache-Control": "no-cache",
                                      "X-Accel-Buffering": "no"})


class NotaDocxRequest(BaseModel):
    fase: str
    variant: int = 0
    projectnaam: str = ""
    markdown: str


@app.post("/api/nota/docx")
def nota_docx(req: NotaDocxRequest):
    """De gestreamde nota-Markdown omzetten naar een opgemaakt Word-document
    (titelpagina, kopstijlen, tabellen, voettekst met paginanummers)."""
    _need_result(req.variant)
    try:
        naam, docx = nota_mod.markdown_naar_docx(
            LAST_RESULT, req.fase, req.variant, req.projectnaam, req.markdown)
    except nota_mod.NotaError as e:
        raise HTTPException(422, str(e))
    return StreamingResponse(
        io.BytesIO(docx), media_type=DOCX_MEDIA,
        headers={"Content-Disposition": f'attachment; filename="{naam}"'})


@app.get("/api/export/nota")
def export_nota(fase: str, variant: int = 0, projectnaam: str = ""):
    """Ontwerpnota in één keer (zonder live weergave): genereren en als
    Word-document teruggeven. Kan 1-3 minuten duren."""
    _need_result(variant)
    try:
        naam, docx = nota_mod.genereer_nota(LAST_RESULT, fase, variant, projectnaam)
    except nota_mod.NotaError as e:
        raise HTTPException(503, str(e))
    return StreamingResponse(
        io.BytesIO(docx), media_type=DOCX_MEDIA,
        headers={"Content-Disposition": f'attachment; filename="{naam}"'})


# ---------------------------------------------------------------------------
# Projecten opslaan / laden
# ---------------------------------------------------------------------------

class ProjectState(BaseModel):
    name: str
    state: dict


def _safe_name(name: str) -> str:
    slug = re.sub(r"[^\w\-]", "_", name.strip())[:60]
    if not slug:
        raise HTTPException(400, "Geef een projectnaam op.")
    return slug


@app.post("/api/project/save")
def project_save(p: ProjectState):
    path = DATA_DIR / f"{_safe_name(p.name)}.json"
    path.write_text(json.dumps({"name": p.name, "state": p.state}, ensure_ascii=False))
    return {"ok": True, "file": path.name}


@app.get("/api/project/list")
def project_list():
    return sorted(f.stem for f in DATA_DIR.glob("*.json"))


@app.get("/api/project/load")
def project_load(name: str):
    global LAST_RESULT
    path = DATA_DIR / f"{_safe_name(name)}.json"
    if not path.exists():
        raise HTTPException(404, "Project niet gevonden.")
    p = json.loads(path.read_text())
    # meegeslagen rekenresultaat weer actief maken, zodat registers,
    # exports en nota's direct werken zonder herberekening
    res = (p.get("state") or {}).get("result")
    if isinstance(res, dict) and res.get("varianten"):
        LAST_RESULT = res
    return p


@app.get("/api/defaults")
def defaults():
    return {"weights": DEFAULT_WEIGHTS, "class_names": CLASS_NAMES,
            "variant_profielen": list(VARIANT_PROFIELEN),
            # optioneel: ingebedde Street View (GOOGLE_MAPS_API_KEY in de
            # omgeving of kabelbed/.env; leeg = fallback via google.com/maps)
            "google_maps_key": os.getenv("GOOGLE_MAPS_API_KEY", "")}


# ---------------------------------------------------------------------------
# ZRO-dossiers: detailpagina, tekening, overeenkomst, bijlagen (FO §6.3)
# ---------------------------------------------------------------------------

class ZroDossierUpdate(BaseModel):
    variant: int = 0
    nr: str
    dossier: dict = Field(default_factory=dict)


class ZroActie(BaseModel):
    variant: int = 0
    nr: str
    achtergrond: str = "dkk"  # dkk | luchtfoto
    projectnaam: str = ""


class ZroBijlageUpload(BaseModel):
    variant: int = 0
    nr: str
    bestandsnaam: str
    data_base64: str


def _zro_item(variant: int, nr: str):
    v = _need_result(variant)
    for item in v["zro"]:
        if item["nr"] == nr:
            return v, item, zro_mod.slug_van(item["perceel"])
    raise HTTPException(404, f"ZRO-item {nr} niet gevonden in deze variant.")


def _zro_detail(item: dict, slug: str, dossier: dict | None = None) -> dict:
    dossier = dossier if dossier is not None else zro_mod.laad_dossier(slug)
    cl = zro_mod.checklist(dossier)
    return {
        "item": item, "slug": slug, "dossier": dossier,
        "checklist": cl, "compleet": all(c["ok"] for c in cl),
        "statussen": zro_mod.ZRO_STATUSSEN, "aard_opties": zro_mod.AARD_OPTIES,
        "vergoeding_voorstel": zro_mod.vergoeding_voorstel(item, dossier),
        "gemeente": LAST_RESULT.get("gemeente") if LAST_RESULT else None,
    }


@app.get("/api/zro/register")
def zro_register(variant: int = 0):
    v = _need_result(variant)
    items = []
    for item in v["zro"]:
        slug = zro_mod.slug_van(item["perceel"])
        d = zro_mod.laad_dossier(slug)
        cl = zro_mod.checklist(d)
        items.append({
            **item, "slug": slug, "status": d["status"],
            "eigenaar_dossier": d["eigenaar_naam"],
            "aard_recht_dossier": d["aard_recht"],
            "vergoeding_eenmalig_eur": d["vergoeding_eenmalig_eur"],
            "vergoeding_jaarlijks_eur": d["vergoeding_jaarlijks_eur"],
            "n_bijlagen": len(d["bijlagen"]),
            "compleet_ok": sum(1 for c in cl if c["ok"]),
            "compleet_totaal": len(cl),
        })
    return {"items": items}


@app.get("/api/zro/detail")
def zro_detail(nr: str, variant: int = 0):
    _v, item, slug = _zro_item(variant, nr)
    return _zro_detail(item, slug)


@app.post("/api/zro/detail")
def zro_detail_save(req: ZroDossierUpdate):
    _v, item, slug = _zro_item(req.variant, req.nr)
    dossier = zro_mod.update_dossier(slug, req.dossier)
    return _zro_detail(item, slug, dossier)


@app.post("/api/zro/tekening")
def zro_tekening(req: ZroActie):
    v, item, slug = _zro_item(req.variant, req.nr)
    if not item.get("geometry"):
        raise HTTPException(422, "Geen perceelgeometrie beschikbaar voor dit ZRO-item.")
    try:
        percelen = pdok.fetch_percelen(tuple(LAST_RESULT["bbox"]))
    except Exception:
        percelen = []
    try:
        pdf = zro_mod.maak_tekening(
            item, v["route"], LAST_RESULT["stations"], percelen,
            LAST_RESULT.get("gemeente"), req.projectnaam,
            achtergrond=req.achtergrond)
    except Exception as e:
        raise HTTPException(500, f"Tekening genereren mislukt: {e}")
    dossier = zro_mod.voeg_bijlage_toe(
        slug, f"ZRO-tekening_{item['nr']}.pdf", pdf, "tekening")
    return _zro_detail(item, slug, dossier)


@app.post("/api/zro/overeenkomst")
def zro_overeenkomst(req: ZroActie):
    _v, item, slug = _zro_item(req.variant, req.nr)
    dossier = zro_mod.laad_dossier(slug)
    ontbreekt = [c["eis"] for c in zro_mod.checklist(dossier) if not c["ok"]]
    if ontbreekt:
        raise HTTPException(422, "Dossier nog niet compleet: " + "; ".join(ontbreekt))
    docx = zro_mod.maak_overeenkomst(item, dossier, LAST_RESULT.get("gemeente"),
                                     req.projectnaam)
    dossier = zro_mod.voeg_bijlage_toe(
        slug, f"ZRO-overeenkomst_{item['nr']}.docx", docx, "overeenkomst")
    # status automatisch bijwerken zolang die nog vóór "overeenkomst opgesteld" staat
    st = zro_mod.ZRO_STATUSSEN
    if dossier["status"] in st and st.index(dossier["status"]) < st.index("overeenkomst opgesteld"):
        dossier["status"] = "overeenkomst opgesteld"
        zro_mod.bewaar_dossier(slug, dossier)
    return _zro_detail(item, slug, dossier)


@app.post("/api/zro/bijlage")
def zro_bijlage_upload(req: ZroBijlageUpload):
    import base64
    _v, item, slug = _zro_item(req.variant, req.nr)
    data = req.data_base64.split(",", 1)[-1]  # eventueel data-URL-prefix strippen
    try:
        inhoud = base64.b64decode(data)
    except Exception:
        raise HTTPException(400, "Bijlage is geen geldige base64-inhoud.")
    if len(inhoud) > 25 * 1024 * 1024:
        raise HTTPException(413, "Bijlage groter dan 25 MB.")
    dossier = zro_mod.voeg_bijlage_toe(slug, req.bestandsnaam, inhoud, "overig")
    return _zro_detail(item, slug, dossier)


@app.delete("/api/zro/bijlage")
def zro_bijlage_delete(slug: str, bestand: str, nr: str = "", variant: int = 0):
    zro_mod.verwijder_bijlage(slug, bestand)
    if nr:
        _v, item, slug2 = _zro_item(variant, nr)
        return _zro_detail(item, slug2)
    return {"ok": True}


@app.get("/api/zro/bijlage/{slug}/{bestand}")
def zro_bijlage_download(slug: str, bestand: str):
    pad = zro_mod.bijlage_pad(slug, bestand)
    if pad is None:
        raise HTTPException(404, "Bijlage niet gevonden.")
    media = {".pdf": "application/pdf",
             ".docx": "application/vnd.openxmlformats-officedocument"
                      ".wordprocessingml.document"}.get(pad.suffix.lower(),
                                                        "application/octet-stream")
    return FileResponse(pad, media_type=media,
                        headers={"Content-Disposition": f'inline; filename="{pad.name}"'})


app.mount("/", StaticFiles(directory=ROOT / "frontend", html=True), name="frontend")
